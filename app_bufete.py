import io
import os
import re
import json
import base64
import shutil
import zipfile
import numpy as np
from datetime import datetime
import streamlit as st
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sklearn.metrics.pairwise import cosine_similarity
from groq import Groq
from markdown_pdf import MarkdownPdf, Section
import gdown

# Restricción de hilos CPU para entornos en la nube (Streamlit Cloud)
os.environ["OMP_NUM_THREADS"] = "2"
os.environ["MKL_NUM_THREADS"] = "2"

# =====================================================================
# INICIALIZACIÓN DE MOTOR OCR LOCAL
# =====================================================================
OCR_LOCAL_DISPONIBLE = False
motor_ocr = None
try:
    from rapidocr_onnxruntime import RapidOCR
    import pypdfium2 as pdfium
    motor_ocr = RapidOCR()
    OCR_LOCAL_DISPONIBLE = True
except Exception:
    OCR_LOCAL_DISPONIBLE = False

st.set_page_config(page_title="Bol-Lex AI - Bufete Digital", page_icon="⚖️", layout="wide")

# =====================================================================
# GESTIÓN DE SESIÓN Y PERSISTENCIA
# =====================================================================
VARIABLES_SESION = {
    "dictamen_actual": "",
    "memorial_actual": "",
    "carpeta_actual": "",
    "codigo_caso_actual": "",
    "titulo_caso_actual": "",
    "materia_detectada": "Civil",
    "historial_consultas": [],
    "evidencia_acumulada": "",
    "borrador_usuario": "",
    "modelo_activo_ia": "",
    "modelos_disponibles_ia": []
}

for key, default in VARIABLES_SESION.items():
    if key not in st.session_state:
        st.session_state[key] = default

def reiniciar_caso():
    for key, default in VARIABLES_SESION.items():
        if key not in ["modelo_activo_ia", "modelos_disponibles_ia"]:
            st.session_state[key] = default
    st.rerun()

def obtener_metadatos_caso(ruta_caso: str) -> dict:
    ruta_meta = os.path.join(ruta_caso, "meta.json")
    if os.path.exists(ruta_meta):
        try:
            with open(ruta_meta, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    nombre_carpeta = os.path.basename(ruta_caso)
    return {"titulo": "", "codigo": nombre_carpeta, "fecha": "", "materia": "Civil"}

def guardar_metadatos_caso(ruta_caso: str, titulo: str, codigo: str, materia: str):
    ruta_meta = os.path.join(ruta_caso, "meta.json")
    datos = {
        "titulo": titulo.strip(),
        "codigo": codigo,
        "materia": materia,
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M")
    }
    with open(ruta_meta, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)

def listar_expedientes_locales() -> dict:
    ruta_base = "Expedientes"
    if not os.path.exists(ruta_base):
        os.makedirs(ruta_base, exist_ok=True)
        return {}
    carpetas = [d for d in os.listdir(ruta_base) if os.path.isdir(os.path.join(ruta_base, d))]
    carpetas.sort(reverse=True)
    mapa = {}
    for c in carpetas:
        ruta_c = os.path.join(ruta_base, c)
        meta = obtener_metadatos_caso(ruta_c)
        nom_corto = meta.get("titulo", "").strip()
        etiqueta = f"{nom_corto} — [{c}]" if (nom_corto and nom_corto != c) else c
        mapa[etiqueta] = c
    return mapa

def cargar_expediente_existente(codigo_caso: str):
    ruta_caso = os.path.join("Expedientes", codigo_caso)
    if not os.path.exists(ruta_caso):
        return
    meta = obtener_metadatos_caso(ruta_caso)
    st.session_state.codigo_caso_actual = codigo_caso
    st.session_state.carpeta_actual = ruta_caso
    st.session_state.titulo_caso_actual = meta.get("titulo", "")
    st.session_state.materia_detectada = meta.get("materia", "Civil")
    
    ruta_dictamen_md = os.path.join(ruta_caso, f"Dictamen_{codigo_caso}.md")
    st.session_state.dictamen_actual = open(ruta_dictamen_md, "r", encoding="utf-8").read() if os.path.exists(ruta_dictamen_md) else ""
    
    ruta_mem_docx = os.path.join(ruta_caso, f"Memorial_{codigo_caso}.docx")
    if os.path.exists(ruta_mem_docx):
        try:
            doc = Document(ruta_mem_docx)
            st.session_state.memorial_actual = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        except Exception:
            st.session_state.memorial_actual = ""
    else:
        st.session_state.memorial_actual = ""
        
    st.session_state.evidencia_acumulada = sincronizar_carpeta_caso(ruta_caso)
    st.rerun()

def eliminar_expediente_actual(codigo_caso: str):
    ruta_caso = os.path.join("Expedientes", codigo_caso)
    if os.path.exists(ruta_caso):
        shutil.rmtree(ruta_caso, ignore_errors=True)
    reiniciar_caso()

# =====================================================================
# GESTIÓN ZIP Y GOOGLE DRIVE
# =====================================================================
def descomprimir_zip(buffer_o_ruta, carpeta_destino: str):
    with zipfile.ZipFile(buffer_o_ruta, 'r') as zip_ref:
        for info in zip_ref.infolist():
            nombre_archivo = os.path.basename(info.filename)
            if not nombre_archivo:
                continue
            if nombre_archivo.lower().endswith((".pdf", ".docx", ".txt")):
                ruta_salida = os.path.join(carpeta_destino, nombre_archivo)
                with zip_ref.open(info) as fuente, open(ruta_salida, "wb") as destino:
                    shutil.copyfileobj(fuente, destino)

def descargar_desde_gdrive(url_o_id: str, carpeta_destino: str) -> bool:
    try:
        if "folders" in url_o_id:
            gdown.download_folder(url=url_o_id, output=carpeta_destino, quiet=False, use_cookies=False)
        else:
            salida = gdown.download(url=url_o_id, output=carpeta_destino + os.sep, quiet=False, fuzzy=True)
            if salida and salida.endswith(".zip"):
                descomprimir_zip(salida, carpeta_destino)
                try:
                    os.remove(salida)
                except Exception:
                    pass
        return True
    except Exception as e:
        st.error(f"Error al descargar desde Google Drive: {e}")
        return False

# =====================================================================
# MOTOR RAG LOCAL DE JURISPRUDENCIA (TCP BOLIVIA)
# =====================================================================
def extraer_datos_pdf_tcp(ruta_pdf: str) -> dict:
    try:
        reader = PdfReader(ruta_pdf)
        texto = "\n".join([p.extract_text() or "" for p in reader.pages])
    except Exception:
        texto = ""
    match_scp = re.search(r"SENTENCIA CONSTITUCIONAL PLURINACIONAL\s+([0-9]{3,4}/[0-9]{4}-[A-Z0-9]+)", texto, re.IGNORECASE)
    scp_numero = f"SCP {match_scp.group(1).strip()}" if match_scp else os.path.basename(ruta_pdf)
    match_sala = re.search(r"(SALA\s+[A-ZÁÉÍÓÚ\s]+)", texto[:2000], re.IGNORECASE)
    sala = match_sala.group(1).strip() if match_sala else "Sala Constitucional"
    patron_ratio = re.search(r"(III\.\s*FUNDAMENTOS\s+JUR[IÍ]DICOS\s+DEL\s+FALLO.*?)(POR\s+TANTO)", texto, re.DOTALL | re.IGNORECASE)
    ratio = patron_ratio.group(1).replace("III. FUNDAMENTOS JURÍDICOS DEL FALLO", "").strip() if patron_ratio else texto[:1500]
    return {"archivo": os.path.basename(ruta_pdf), "scp_numero": scp_numero, "sala": sala, "ratio_decidendi": ratio}

class MotorRAGJurisprudencial:
    def __init__(self, carpeta: str = "sentencias_pdf"):
        self.docs = []
        if os.path.exists(carpeta):
            self.docs = [extraer_datos_pdf_tcp(os.path.join(carpeta, arch)) for arch in os.listdir(carpeta) if arch.endswith(".pdf")]
        if not self.docs:
            self.docs = [{
                "archivo": "SCP_0084_2017_S3.pdf",
                "scp_numero": "SCP 0084/2017-S3",
                "sala": "Sala Tercera",
                "ratio_decidendi": "Tutela directa del fuero laboral y conminatoria de reincorporación."
            }]
        try:
            from sentence_transformers import SentenceTransformer
            self.encoder = SentenceTransformer("paraphrase-multilingual-MiniLM-L12-v2")
            self.vectores = self.encoder.encode([d["ratio_decidendi"] for d in self.docs])
        except Exception:
            self.encoder = None
            self.vectores = None
    
    def buscar(self, consulta: str) -> dict:
        if not self.docs or self.encoder is None or self.vectores is None:
            return self.docs[0] if self.docs else {"scp_numero": "S/N", "sala": "-", "ratio_decidendi": "Sin jurisprudencia local indexada."}
        v_q = self.encoder.encode([consulta])
        sims = cosine_similarity(v_q, self.vectores)[0]
        idx = int(np.argmax(sims))
        res = self.docs[idx].copy()
        res["score"] = float(sims[idx])
        return res

@st.cache_resource
def cargar_motor_rag():
    return MotorRAGJurisprudencial("sentencias_pdf")

# =====================================================================
# MOTOR DE IA RESILIENTE (GROQ) CON AUTO-DETECCIÓN Y FALLBACK
# =====================================================================
@st.cache_resource
def inicializar_motor_ia():
    api_key = os.environ.get("GROQ_API_KEY", "")
    if not api_key:
        try:
            api_key = st.secrets["GROQ_API_KEY"]
        except Exception:
            api_key = ""
    cliente = Groq(api_key=api_key)

    # Candidatos ordenados por capacidad de razonamiento jurídico y velocidad
    candidatos_base = [
        "llama-3.3-70b-versatile",
        "llama-3.1-70b-versatile",
        "llama-3.1-8b-instant",
        "llama3-70b-8192",
        "llama3-8b-8192",
        "mixtral-8x7b-32768",
        "gemma2-9b-it",
        "qwen/qwen3.8-27b",
    ]

    ids_cuenta = []
    try:
        modelos = cliente.models.list()
        ids_cuenta = [
            m.id for m in modelos.data 
            if not any(k in m.id.lower() for k in ["guard", "whisper", "orpheus", "vision"])
        ]
    except Exception:
        pass

    # Combinación priorizada respetando la respuesta en vivo de la cuenta
    lista_candidatos = [c for c in candidatos_base if c in ids_cuenta]
    for m in ids_cuenta:
        if m not in lista_candidatos:
            lista_candidatos.append(m)
    for c in candidatos_base:
        if c not in lista_candidatos:
            lista_candidatos.append(c)

    # Prueba de fuego en caliente (dry-run): seleccionar el primer modelo con HTTP 200
    modelo_confirmado = "llama-3.1-8b-instant"
    for cand in lista_candidatos:
        try:
            cliente.chat.completions.create(
                model=cand,
                messages=[{"role": "user", "content": "ping"}],
                max_tokens=2
            )
            modelo_confirmado = cand
            break
        except Exception:
            continue

    return cliente, modelo_confirmado, lista_candidatos

cliente_groq, modelo_predeterminado, lista_modelos_disponibles = inicializar_motor_ia()

if not st.session_state.modelo_activo_ia:
    st.session_state.modelo_activo_ia = modelo_predeterminado
st.session_state.modelos_disponibles_ia = lista_modelos_disponibles

def invocar_groq_con_fallback(mensajes: list, temperatura: float = 0.05, max_tokens: int = None) -> str:
    """
    Ejecuta la inferencia contra Groq con manejo defensivo y conmutación automática
    ante errores 404/model_not_found o cuotas excedidas.
    """
    candidatos = st.session_state.get("modelos_disponibles_ia", [st.session_state.modelo_activo_ia])
    modelo_actual = st.session_state.get("modelo_activo_ia", modelo_predeterminado)
    
    intentos = [modelo_actual] + [m for m in candidatos if m != modelo_actual]
    ultimo_error = None

    for mod in intentos:
        try:
            args = {
                "model": mod,
                "messages": mensajes,
                "temperature": temperatura
            }
            if max_tokens:
                args["max_tokens"] = max_tokens
            res = cliente_groq.chat.completions.create(**args)
            if mod != st.session_state.modelo_activo_ia:
                st.session_state.modelo_activo_ia = mod
            return res.choices[0].message.content
        except Exception as e:
            err_str = str(e).lower()
            ultimo_error = e
            if any(k in err_str for k in ["model_not_found", "does not exist", "404", "rate_limit"]):
                continue
            raise e

    if ultimo_error:
        raise ultimo_error
    return ""

# =====================================================================
# EXTRACTOR DOCUMENTAL (TEXTO NATIVO + OCR LOCAL DEFENSIVO)
# =====================================================================
def ocr_pagina_local(imagen_rgb_array) -> str:
    """Aplica OCR local mediante RapidOCR ONNX."""
    if not OCR_LOCAL_DISPONIBLE or motor_ocr is None:
        return ""
    try:
        resultado, _ = motor_ocr(imagen_rgb_array)
        if resultado:
            return "\n".join([linea[1] for linea in resultado if linea and len(linea) > 1])
    except Exception:
        pass
    return ""

def extraer_texto_archivo(ruta_o_buffer, nombre: str, carpeta_caso: str = None) -> str:
    nl = nombre.lower()
    try:
        if nl.endswith(".txt"):
            if hasattr(ruta_o_buffer, "read"):
                return ruta_o_buffer.read().decode("utf-8", errors="ignore")
            with open(ruta_o_buffer, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif nl.endswith(".docx"):
            doc = Document(ruta_o_buffer)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        elif nl.endswith(".pdf"):
            if carpeta_caso:
                ruta_cache = os.path.join(carpeta_caso, f"{nombre}.cache.txt")
                if os.path.exists(ruta_cache):
                    with open(ruta_cache, "r", encoding="utf-8") as fc:
                        return fc.read()

            reader = PdfReader(ruta_o_buffer)
            paginas = [p.extract_text() or "" for p in reader.pages]
            texto_digital = "\n".join(paginas).strip()
            
            # Si el PDF tiene texto digital suficiente (>80 caracteres), utilizarlo
            if len(texto_digital) > 80:
                if carpeta_caso:
                    with open(os.path.join(carpeta_caso, f"{nombre}.cache.txt"), "w", encoding="utf-8") as fc:
                        fc.write(texto_digital)
                return texto_digital

            # Si es un escaneo físico sin capa de texto, aplicar OCR local
            if OCR_LOCAL_DISPONIBLE:
                st.toast(f"🔍 OCR Local en: {nombre}...")
                try:
                    if isinstance(ruta_o_buffer, str):
                        doc_pdf = pdfium.PdfDocument(ruta_o_buffer)
                    else:
                        if hasattr(ruta_o_buffer, "seek"):
                            ruta_o_buffer.seek(0)
                        doc_pdf = pdfium.PdfDocument(ruta_o_buffer.read())
                    
                    texto_ocr = []
                    limite_pags = min(len(doc_pdf), 6)
                    for idx in range(limite_pags):
                        page = doc_pdf[idx]
                        bitmap = page.render(scale=1.4)
                        pil_img = bitmap.to_pil().convert("RGB")
                        img_array = np.array(pil_img)
                        txt_p = ocr_pagina_local(img_array)
                        if txt_p:
                            texto_ocr.append(f"[Página {idx+1}]\n{txt_p}")
                    
                    texto_completo = "\n".join(texto_ocr)
                    if texto_completo.strip():
                        if carpeta_caso:
                            with open(os.path.join(carpeta_caso, f"{nombre}.cache.txt"), "w", encoding="utf-8") as fc:
                                fc.write(texto_completo)
                        return texto_completo
                except Exception:
                    pass

            return f"[Documento {nombre}: Documento escaneado o sin texto reconocible]"
        return ""
    except Exception as e:
        return f"Error procesando {nombre}: {e}"

def sincronizar_carpeta_caso(ruta_carpeta: str) -> str:
    acumulado = ""
    if os.path.exists(ruta_carpeta):
        archivos = [
            a for a in os.listdir(ruta_carpeta) 
            if a.endswith((".pdf", ".txt", ".docx")) 
            and not a.startswith(("Dictamen_", "Memorial_"))
            and not a.endswith((".cache.txt", ".resumen.json"))
        ]
        for arch in archivos:
            ruta_completa = os.path.join(ruta_carpeta, arch)
            texto = extraer_texto_archivo(ruta_completa, arch, carpeta_caso=ruta_carpeta)
            acumulado += f"\n--- DOCUMENTO: {arch} ---\n{texto[:4000].strip()}\n"
    return acumulado

# =====================================================================
# MAP-REDUCE FORENSE (EXTRACCIÓN DOCUMENTAL MASIVA)
# =====================================================================
def resumir_documento_individual(nombre_archivo: str, texto: str, carpeta_caso: str) -> dict:
    ruta_resumen = os.path.join(carpeta_caso, f"{nombre_archivo}.resumen.json")
    if os.path.exists(ruta_resumen):
        try:
            with open(ruta_resumen, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass

    prompt_extract = f"""
Analiza este documento del expediente judicial/contractual boliviano y extrae solo hechos y datos fácticos:
DOCUMENTO: {nombre_archivo}
CONTENIDO:
{texto[:5000]}

Responde en este formato conciso y estricto:
- Tipo de documento y fecha:
- Partes identificadas (Acreedor / Deudor exactos, NIT, personas):
- Objeto del contrato / trámite / proceso:
- Montos económicos y saldos (con moneda exacta y comprobantes citados):
- Cumplimiento, mora o controversia:
REGLA: Si un dato no figura, escribe 'No especificado'. No inventes información.
"""
    try:
        sintesis = invocar_groq_con_fallback(
            mensajes=[{"role": "user", "content": prompt_extract}],
            temperatura=0.0
        ).strip()
        datos = {"archivo": nombre_archivo, "sintesis": sintesis}
        with open(ruta_resumen, "w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False, indent=2)
        return datos
    except Exception:
        return {"archivo": nombre_archivo, "sintesis": texto[:600]}

def construir_expediente_estructurado(carpeta_caso: str) -> str:
    if not os.path.exists(carpeta_caso):
        return ""
    archivos = [
        a for a in os.listdir(carpeta_caso) 
        if a.endswith((".pdf", ".txt", ".docx")) 
        and not a.startswith(("Dictamen_", "Memorial_"))
        and not a.endswith((".cache.txt", ".resumen.json"))
    ]
    if not archivos:
        return ""
    
    informe_consolidado = []
    for arch in archivos:
        ruta_arch = os.path.join(carpeta_caso, arch)
        texto = extraer_texto_archivo(ruta_arch, arch, carpeta_caso=carpeta_caso)
        sintesis = resumir_documento_individual(arch, texto, carpeta_caso)
        informe_consolidado.append(f"### DOCUMENTO: {arch}\n{sintesis.get('sintesis', '')}\n")
        
    return "\n".join(informe_consolidado)

# =====================================================================
# GENERACIÓN DE DICTAMEN INTEGRAL
# =====================================================================
def ejecutar_dictamen_integral(hechos: str, docs_estructurados: str, jurisprudencia: dict = None) -> str:
    prompt_sistema = """
Actúas como Consultor Jurídico Senior en Bolivia en Derecho Civil, Administrativo y Contrataciones del Estado (Ley 1178, Ley 439 y DS 0181).
Elabora un DICTAMEN LEGAL CONTUNDENTE Y RIGUROSO.

ESTRUCTURA DEL DICTAMEN:
1. IDENTIFICACIÓN DE LAS PARTES (Acreedor, deudor, personerías, NIT).
2. ANTECEDENTES Y RELACIÓN CONTRACTUAL (Cronología exacta de contratos, adendas, actas de recepción definitiva y notas).
3. DETERMINACIÓN DE LA OBLIGACIÓN ECONÓMICA (Monto total contractual, pagos parciales acreditados y saldo impago exigible).
4. PRECEDENTES JURISPRUDENCIALES Y MARCO CONSTITUCIONAL (Doctrina aplicable del TCP o normativa boliviana).
5. VIABILIDAD PROCESAL Y ESTRATEGIA DE LITIGIO (Vía ejecutiva/coactiva, contencioso-administrativa, excepciones y medidas cautelares).
6. PLAN DE ACCIÓN INMEDIATO (Requerimientos, agotamiento de vía y plazos fatales).

REGLA DE ORO:
- Basa tus conclusiones ÚNICAMENTE en los datos reales del expediente.
- Si no hay comprobante bancario de pago de un saldo, la deuda permanece vigente y exigible.
"""
    prompt_usuario = f"""
INSTRUCCIONES Y ANTECEDENTES DEL CLIENTE:
{hechos}

PRECEDENTE CONSTITUCIONAL VINCULANTE:
- Fallo: {jurisprudencia.get('scp_numero', 'Sin precedente específico')} ({jurisprudencia.get('sala', '')})
- Ratio Decidendi: {jurisprudencia.get('ratio_decidendi', 'Aplicación de normativa general')}

EXPEDIENTE DOCUMENTADO:
{docs_estructurados if docs_estructurados else 'Sin documentos adjuntos adicionales.'}
"""
    return invocar_groq_con_fallback(
        mensajes=[{"role": "system", "content": prompt_sistema}, {"role": "user", "content": prompt_usuario}],
        temperatura=0.05
    )

# =====================================================================
# REDACTOR FORENSE DE MEMORIALES (WORD .DOCX Y PDF)
# =====================================================================
def redactar_memorial_forense(hechos: str, docs: str, borrador_usuario: str, materia: str, abogado: str) -> str:
    prompt_sistema = f"""
Actúas como el Abogado Litigante patrocinante ({abogado}) en el Estado Plurinacional de Bolivia.
Redacta un MEMORIAL FORENSE COMPLETO con formato judicial boliviano riguroso:
- SEÑOR JUEZ (Autoridad judicial competente en letras mayúsculas)
- SUMA (Petitorio en mayúsculas negritas)
- GENERALES DE LEY DEL DEMANDANTE / RECURRENTE
- ANTECEDENTES FÁCTICOS Y RELACIÓN DE HECHOS
- FUNDAMENTACIÓN DE DERECHO (Códigos y Leyes de Bolivia)
- PETITORIO CONCRETO Y EXPEDITO
- OTROSÍES (Otrosí 1° Pruebas, Otrosí 2° Domicilio procesal, Otrosí 3° Honorarios profesionales).
Materia procesal: {materia}.
"""
    prompt_usuario = f"""
INSTRUCCIONES DEL CASO:
{hechos}

EVIDENCIA DOCUMENTAL DEL EXPEDIENTE:
{docs[:7000]}

BORRADOR O NOTAS ADICIONALES DEL ABOGADO:
{borrador_usuario if borrador_usuario else 'Redactar memorial formal íntegro.'}
"""
    return invocar_groq_con_fallback(
        mensajes=[{"role": "system", "content": prompt_sistema}, {"role": "user", "content": prompt_usuario}],
        temperatura=0.1
    )

def exportar_memorial_docx(texto_md: str, ruta_salida: str):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.3)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    for linea in texto_md.split("\n"):
        l = linea.strip()
        if not l:
            continue
        if l.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(l.replace("# ", ""))
            run.bold = True
            run.font.size = Pt(14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif l.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(l.replace("## ", ""))
            run.bold = True
            run.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(10)
        elif l.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(l.replace("### ", ""))
            run.bold = True
        elif l.startswith("SEÑOR JUEZ") or l.startswith("SUMA:"):
            p = doc.add_paragraph()
            run = p.add_run(l.replace("**", ""))
            run.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
        else:
            p = doc.add_paragraph()
            texto_limpio = l.replace("**", "").replace("*", "")
            p.add_run(texto_limpio)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
    doc.save(ruta_salida)

def relacionar_todos_los_casos() -> str:
    mapa = listar_expedientes_locales()
    if not mapa:
        return "No hay expedientes registrados para relacionar."
    
    resumen_global = ""
    for etiq, cod in mapa.items():
        ruta = os.path.join("Expedientes", cod)
        ruta_dict = os.path.join(ruta, f"Dictamen_{cod}.md")
        resumen_dict = "Sin dictamen generado aún."
        if os.path.exists(ruta_dict):
            with open(ruta_dict, "r", encoding="utf-8") as f:
                resumen_dict = f.read()[:1000]
        resumen_global += f"\n=== EXPEDIENTE: {etiq} ===\n{resumen_dict}\n"

    prompt = f"""
Actúas como Director Estratégico del Bufete.
Analiza la totalidad de expedientes del despacho e identifica conexidades, deudores recurrentes, NITs comunes o riesgos procesales cruzados:
{resumen_global}
"""
    return invocar_groq_con_fallback(
        mensajes=[{"role": "user", "content": prompt}],
        temperatura=0.1
    )

# =====================================================================
# INTERFAZ STREAMLIT (CONSOLIDACIÓN PMV)
# =====================================================================
st.title("⚖️ Bol-Lex AI: Bufete Digital Litigante")

motor_rag = cargar_motor_rag()

# BARRA LATERAL
with st.sidebar:
    st.success(f"🤖 IA Activa: `{st.session_state.modelo_activo_ia}`")
    st.info(f"👁️ Motor OCR: `{'Local ONNX Activo' if OCR_LOCAL_DISPONIBLE else 'PyPDF Digital'}`")
    
    # Selector de modelo alternativo
    opciones_modelos = st.session_state.get("modelos_disponibles_ia", [st.session_state.modelo_activo_ia])
    idx_mod = opciones_modelos.index(st.session_state.modelo_activo_ia) if st.session_state.modelo_activo_ia in opciones_modelos else 0
    modelo_elegido = st.selectbox("Cambiar Modelo de IA:", options=opciones_modelos, index=idx_mod)
    if modelo_elegido != st.session_state.modelo_activo_ia:
        st.session_state.modelo_activo_ia = modelo_elegido
        st.rerun()

    st.markdown("---")
    st.subheader("⚡ Gestión del Expediente")
    col_g1, col_g2 = st.columns(2)
    with col_g1:
        if st.button("💾 Guardar Caso", use_container_width=True):
            if st.session_state.carpeta_actual and os.path.exists(st.session_state.carpeta_actual):
                guardar_metadatos_caso(
                    st.session_state.carpeta_actual,
                    st.session_state.titulo_caso_actual,
                    st.session_state.codigo_caso_actual,
                    st.session_state.materia_detectada
                )
                if st.session_state.dictamen_actual:
                    with open(os.path.join(st.session_state.carpeta_actual, f"Dictamen_{st.session_state.codigo_caso_actual}.md"), "w", encoding="utf-8") as f:
                        f.write(st.session_state.dictamen_actual)
                st.toast("✅ Guardado con éxito.")
                st.rerun()
    with col_g2:
        if st.button("🔄 Nuevo Caso", type="secondary", use_container_width=True):
            reiniciar_caso()

    mapa_expedientes = listar_expedientes_locales()
    st.markdown("##### 📂 Historial de Expedientes")
    opciones = ["-- Seleccionar caso --"] + list(mapa_expedientes.keys())
    idx_defecto = 0
    for i, op in enumerate(opciones):
        if op != "-- Seleccionar caso --" and mapa_expedientes[op] == st.session_state.codigo_caso_actual:
            idx_defecto = i
            break
            
    seleccion = st.selectbox("Expedientes guardados:", options=opciones, index=idx_defecto)
    if seleccion != "-- Seleccionar caso --":
        cod_elegido = mapa_expedientes[seleccion]
        if cod_elegido != st.session_state.codigo_caso_actual:
            if st.button("📥 Cargar este Caso", use_container_width=True):
                cargar_expediente_existente(cod_elegido)

    if st.session_state.codigo_caso_actual:
        st.markdown("---")
        with st.expander("🗑️ Zona de Peligro"):
            if st.button("❌ Borrar este expediente", type="primary", use_container_width=True):
                eliminar_expediente_actual(st.session_state.codigo_caso_actual)

    st.markdown("---")
    st.subheader("📁 Archivos del Caso Activo")
    archivos_adjuntos = st.file_uploader(
        "Subir PDF, Word (.docx), TXT o comprimidos (.ZIP)",
        type=["pdf", "txt", "docx", "zip"],
        accept_multiple_files=True
    )
    
    url_gdrive = st.text_input("🔗 O pegar enlace de Google Drive:")
    if url_gdrive and st.button("📥 Importar desde Drive", use_container_width=True):
        cod_caso = st.session_state.codigo_caso_actual or f"BOLLEX-{datetime.now().strftime('%Y%m%d-%H%M')}"
        ruta_dest = os.path.join("Expedientes", cod_caso)
        os.makedirs(ruta_dest, exist_ok=True)
        st.session_state.carpeta_actual = ruta_dest
        st.session_state.codigo_caso_actual = cod_caso
        with st.spinner("Descargando desde Google Drive..."):
            if descargar_desde_gdrive(url_gdrive, ruta_dest):
                st.session_state.evidencia_acumulada = sincronizar_carpeta_caso(ruta_dest)
                st.toast("✅ Archivos de Drive importados.")
                st.rerun()

    if st.session_state.carpeta_actual and os.path.exists(st.session_state.carpeta_actual):
        st.markdown(f"**Expediente:** `{st.session_state.codigo_caso_actual}`")
        st.caption("Documentos en este expediente:")
        for f in os.listdir(st.session_state.carpeta_actual):
            if not f.endswith((".cache.txt", ".resumen.json", "meta.json")):
                st.text(f"• {f}")

# PESTAÑAS PRINCIPALES DEL PMV
tab1, tab2, tab3, tab4 = st.tabs([
    "📋 Dictamen y Diagnóstico", 
    "📝 Redactor Forense (.docx)", 
    "💬 Consultorio Jurídico", 
    "🗄️ Archivo Documental"
])

# =====================================================================
# TAB 1: DICTAMEN Y DIAGNÓSTICO JURÍDICO
# =====================================================================
with tab1:
    st.subheader("📋 Dictamen y Diagnóstico Jurídico Estratégico")
    col_t1, col_t2 = st.columns([2, 1])
    with col_t1:
        titulo_ingresado = st.text_input(
            "🏷️ Identificador del caso:", 
            value=st.session_state.titulo_caso_actual, 
            placeholder="Ej: Caso Deuda EBC a Ecotraffic"
        )
        if titulo_ingresado != st.session_state.titulo_caso_actual:
            st.session_state.titulo_caso_actual = titulo_ingresado
    with col_t2:
        materia_sel = st.selectbox(
            "⚖️ Materia Principal:",
            options=["Civil", "Laboral", "Constitucional", "Penal", "Familia", "Administrativo"],
            index=["Civil", "Laboral", "Constitucional", "Penal", "Familia", "Administrativo"].index(st.session_state.materia_detectada) if st.session_state.materia_detectada in ["Civil", "Laboral", "Constitucional", "Penal", "Familia", "Administrativo"] else 0
        )
        st.session_state.materia_detectada = materia_sel

    caso_cliente = st.text_area(
        "✍️ Planteamiento del caso e instrucciones para el bufete:",
        value="Revisa los antecedentes del caso adjunto y establece estrategia jurídica y acciones procesales.",
        height=90
    )

    if st.button("🚀 Analizar Caso y Generar Dictamen Estratégico", type="primary"):
        codigo_caso = st.session_state.codigo_caso_actual or f"BOLLEX-{datetime.now().strftime('%Y%m%d-%H%M')}"
        ruta_carpeta = os.path.join("Expedientes", codigo_caso)
        os.makedirs(ruta_carpeta, exist_ok=True)
        st.session_state.carpeta_actual = ruta_carpeta
        st.session_state.codigo_caso_actual = codigo_caso

        guardar_metadatos_caso(
            ruta_carpeta, 
            st.session_state.titulo_caso_actual or codigo_caso, 
            codigo_caso, 
            st.session_state.materia_detectada
        )

        if archivos_adjuntos:
            for arch in archivos_adjuntos:
                if arch.name.lower().endswith(".zip"):
                    descomprimir_zip(arch, ruta_carpeta)
                else:
                    with open(os.path.join(ruta_carpeta, arch.name), "wb") as f:
                        f.write(arch.getbuffer())

        with st.spinner("Procesando expediente con Map-Reduce documental e IA defensiva..."):
            try:
                # 1. Map-Reduce de evidencias
                docs_jerarquizados = construir_expediente_estructurado(ruta_carpeta)
                st.session_state.evidencia_acumulada = docs_jerarquizados

                # 2. Búsqueda de jurisprudencia en el TCP
                precedente = motor_rag.buscar(caso_cliente)

                # 3. Generación del dictamen legal
                dictamen = ejecutar_dictamen_integral(caso_cliente, docs_jerarquizados, jurisprudencia=precedente)
                st.session_state.dictamen_actual = dictamen

                # 4. Persistencia en Markdown y PDF
                with open(os.path.join(ruta_carpeta, f"Dictamen_{codigo_caso}.md"), "w", encoding="utf-8") as f:
                    f.write(dictamen)

                pdf = MarkdownPdf(toc_level=0)
                pdf.add_section(Section(dictamen))
                pdf.save(os.path.join(ruta_carpeta, f"Dictamen_{codigo_caso}.pdf"))
                st.success("✅ Dictamen concluido con éxito.")
                st.rerun()
            except Exception as e:
                st.error(f"❌ Error en procesamiento: {e}")

    if st.session_state.dictamen_actual:
        st.divider()
        col1, col2 = st.columns([3, 1])
        with col1:
            st.subheader(f"📄 Dictamen: {st.session_state.titulo_caso_actual or st.session_state.codigo_caso_actual}")
        with col2:
            ruta_pdf = os.path.join(st.session_state.carpeta_actual, f"Dictamen_{st.session_state.codigo_caso_actual}.pdf")
            if os.path.exists(ruta_pdf):
                with open(ruta_pdf, "rb") as fp:
                    st.download_button(
                        "📥 Descargar Dictamen en PDF", 
                        fp, 
                        file_name=f"Dictamen_{st.session_state.codigo_caso_actual}.pdf", 
                        mime="application/pdf", 
                        use_container_width=True
                    )
        st.markdown(st.session_state.dictamen_actual)

# =====================================================================
# TAB 2: REDACTOR FORENSE DE MEMORIALES
# =====================================================================
with tab2:
    st.subheader("📝 Redactor Forense de Memoriales Judiciales (Word .docx y PDF)")
    col_ab1, col_ab2 = st.columns([2, 1])
    with col_ab1:
        abogado_firmante = st.text_input("Abogado patrocinante:", value="Dr. Rolando Gutierrez - M.C.A. 45892")
    with col_ab2:
        tipo_accion = st.text_input("Tipo de memorial / Suma:", value="Demanda Ejecutiva / Acción Tutelar")

    borrador_subido = st.file_uploader("Adjuntar borrador o petitorio complementario (.txt o .docx):", type=["txt", "docx"], key="upload_borrador")
    texto_borrador_manual = st.text_area("Notas e instrucciones específicas para la redacción:", height=100)

    if st.button("⚖️ Redactar Memorial Oficial en Formato Forense", type="primary"):
        with st.spinner("Redactando memorial forense según normativa boliviana..."):
            try:
                borrador_final = ""
                if borrador_subido:
                    borrador_final += f"\n--- BORRADOR ADJUNTO ---\n{extraer_texto_archivo(borrador_subido, borrador_subido.name, st.session_state.carpeta_actual)}\n"
                if texto_borrador_manual:
                    borrador_final += f"\n--- NOTAS ADICIONALES ---\n{texto_borrador_manual}\n"

                mem_texto = redactar_memorial_forense(
                    hechos=caso_cliente + "\nSuma pretendida: " + tipo_accion,
                    docs=st.session_state.evidencia_acumulada,
                    borrador_usuario=borrador_final,
                    materia=st.session_state.materia_detectada,
                    abogado=abogado_firmante
                )
                st.session_state.memorial_actual = mem_texto

                if st.session_state.carpeta_actual:
                    ruta_mem_docx = os.path.join(st.session_state.carpeta_actual, f"Memorial_{st.session_state.codigo_caso_actual}.docx")
                    exportar_memorial_docx(mem_texto, ruta_mem_docx)
                    
                    # Generar también PDF del memorial
                    pdf_mem = MarkdownPdf(toc_level=0)
                    pdf_mem.add_section(Section(mem_texto))
                    pdf_mem.save(os.path.join(st.session_state.carpeta_actual, f"Memorial_{st.session_state.codigo_caso_actual}.pdf"))

                st.success("✅ Memorial redactado y compilado en .docx y .pdf.")
                st.rerun()
            except Exception as e:
                st.error(f"❌ Error redactando memorial: {e}")

    if st.session_state.memorial_actual:
        st.divider()
        col_m1, col_m2, col_m3 = st.columns([2, 1, 1])
        with col_m1:
            st.subheader("📄 Memorial Forense Generado")
        with col_m2:
            ruta_mem_docx = os.path.join(st.session_state.carpeta_actual, f"Memorial_{st.session_state.codigo_caso_actual}.docx")
            if os.path.exists(ruta_mem_docx):
                with open(ruta_mem_docx, "rb") as f_docx:
                    st.download_button(
                        label="📥 Descargar Word (.docx)",
                        data=f_docx,
                        file_name=f"Memorial_{st.session_state.codigo_caso_actual}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
        with col_m3:
            ruta_mem_pdf = os.path.join(st.session_state.carpeta_actual, f"Memorial_{st.session_state.codigo_caso_actual}.pdf")
            if os.path.exists(ruta_mem_pdf):
                with open(ruta_mem_pdf, "rb") as f_mpdf:
                    st.download_button(
                        label="📥 Descargar PDF",
                        data=f_mpdf,
                        file_name=f"Memorial_{st.session_state.codigo_caso_actual}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )
        st.markdown(st.session_state.memorial_actual)

# =====================================================================
# TAB 3: CONSULTORIO JURÍDICO Y RELACIÓN INTERCASOS
# =====================================================================
with tab3:
    st.subheader("💬 Consultorio Jurídico y Relación Intercasos")
    
    st.markdown("#### 🔗 Informe del Director: Relación y Conexidad Global")
    if st.button("🌐 Analizar Relación de Todos los Casos del Bufete", type="secondary"):
        with st.spinner("Analizando conexidad entre todos los expedientes del bufete..."):
            try:
                informe_cruzado = relacionar_todos_los_casos()
                st.markdown(informe_cruzado)
                st.divider()
            except Exception as e:
                st.error(f"Error generando análisis global: {e}")

    st.markdown("#### 🗨️ Consultas sobre el Expediente Activo")
    for preg, resp in st.session_state.historial_consultas:
        with st.chat_message("user"):
            st.write(preg)
        with st.chat_message("assistant"):
            st.markdown(resp)

    pregunta = st.chat_input("Consulta jurídica sobre los antecedentes o estrategia de este caso...")
    if pregunta:
        with st.chat_message("user"):
            st.write(pregunta)
        
        prompt_chat = f"""
Actúas como consultor jurídico litigante.
Dictamen elaborado:
{st.session_state.dictamen_actual[:2500]}

Evidencias documentadas:
{st.session_state.evidencia_acumulada[:6000]}

Pregunta del abogado:
{pregunta}
"""
        try:
            resp_texto = invocar_groq_con_fallback(
                mensajes=[{"role": "user", "content": prompt_chat}],
                temperatura=0.1
            )
            with st.chat_message("assistant"):
                st.markdown(resp_texto)
            st.session_state.historial_consultas.append((pregunta, resp_texto))
        except Exception as e:
            st.error(f"Error en consulta: {e}")

# =====================================================================
# TAB 4: ARCHIVO DOCUMENTAL DEL EXPEDIENTE
# =====================================================================
with tab4:
    st.subheader("🗄️ Archivo Documental del Expediente Activo")
    if not st.session_state.carpeta_actual or not os.path.exists(st.session_state.carpeta_actual):
        st.info("No hay un expediente activo cargado.")
    else:
        archivos_disco = [f for f in os.listdir(st.session_state.carpeta_actual) if not f.endswith((".cache.txt", ".resumen.json", "meta.json"))]
        
        col_d1, col_d2 = st.columns(2)
        with col_d1:
            st.markdown("#### 📑 Documentos Elaborados por Bol-Lex")
            elaborados = [f for f in archivos_disco if f.startswith(("Dictamen_", "Memorial_"))]
            if elaborados:
                for arch in elaborados:
                    ruta_arch = os.path.join(st.session_state.carpeta_actual, arch)
                    with open(ruta_arch, "rb") as f_el:
                        st.download_button(
                            label=f"⬇️ Descargar {arch}",
                            data=f_el,
                            file_name=arch,
                            key=f"dl_el_{arch}"
                        )
            else:
                st.caption("Aún no se han generado dictámenes o memoriales para este expediente.")

        with col_d2:
            st.markdown("#### 📎 Documentos y Pruebas Adjuntas")
            adjuntos = [f for f in archivos_disco if not f.startswith(("Dictamen_", "Memorial_"))]
            if adjuntos:
                for arch in adjuntos:
                    ruta_arch = os.path.join(st.session_state.carpeta_actual, arch)
                    with open(ruta_arch, "rb") as f_adj:
                        st.download_button(
                            label=f"📎 Descargar {arch}",
                            data=f_adj,
                            file_name=arch,
                            key=f"dl_adj_{arch}"
                        )
            else:
                st.caption("No hay documentos adjuntos en la carpeta.")