"""
Agente 4: Redactor Forense de Memoriales Judiciales (Word .docx y PDF).
Redacta escritos judiciales formales según los usos y prácticas forenses del foro boliviano.
"""
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from markdown_pdf import MarkdownPdf, Section

from ..llm_client import MotorLLM

class AgenteRedactorForense:
    def __init__(self, motor_llm: MotorLLM):
        self.llm = motor_llm

    def redactar_memorial(
        self,
        titulo_caso: str,
        materia: str,
        suma_pretendida: str,
        abogado_firmante: str,
        evidencia_auditada: str,
        instrucciones_especificas: str = ""
    ) -> str:
        """
        Redacta el memorial judicial íntegro listo para litigio.
        """
        prompt_sistema = f"""
Actúas como el Abogado Litigante Redactor Forense Senior en Bolivia ({abogado_firmante}).
Tu misión es redactar un MEMORIAL JUDICIAL OFICIAL FORMAL rigurosamente apegado a la técnica forense boliviana.
Materia: {materia}.

ESTRUCTURA OBLIGATORIA DEL MEMORIAL BOLIVIANO:
# SEÑOR JUEZ PÚBLICO EN MATERIA CIVIL Y COMERCIAL DE TURNO DE LA CAPITAL (o Sala Constitucional de Turno)

## SUMA: {suma_pretendida.upper()}

## I. GENERALES DE LEY DEL DEMANDANTE / RECURRENTE
- Nombres, apellidos, cédula de identidad, NIT, domicilio real y procesal. En caso de persona jurídica, acreditar personería mediante Poder Notarial.

## II. IDENTIFICACIÓN DE LA PARTE DEMANDADA / RECURRIDA
- Nombre exacto de la persona natural o jurídica demandada (ej. Empresa Estratégica Boliviana de Construcción y Conservación de Infraestructura Civil - EBC), NIT y datos de su representante legal o MAE.

## III. ANTECEDENTES FÁCTICOS DE LA CONTROVERSIA
- Exposición circunstanciada, clara y cronológica de los contratos, actas de recepción, planillas impagas y cartas notariales de conminatoria.

## IV. FUNDAMENTACIÓN DE DERECHO
- Citas expresas de los artículos aplicables del Código Civil, Código Procesal Civil (Ley 439) o Código Procesal Constitucional (Ley 254).
- Cita de la Ratio Decidendi de la Sentencia Constitucional Plurinacional vinculante.

## V. PETITORIO
- Pronunciamiento expreso solicitando la admisión de la demanda/recurso, intimación de pago dentro de tercero día, o concesión de la tutela solicitada.

## OTROSÍ PRIMERO.- (DE LA PRUEBA DOCUMENTAL)
- Detalle enumerado de las pruebas preconstituidas adjuntas (Contratos, actas, cartas notariales).

## OTROSÍ SEGUNDO.- (DE LAS MEDIDAS CAUTELARES)
- Solicitud de embargo preventivo o retención de fondos ante la Autoridad de Supervisión del Sistema Financiero (ASFI).

## OTROSÍ TERCERO.- (HONORARIOS PROFESIONALES Y DOMICILIO)
- Arancel del Colegio de Abogados de Bolivia y señalamiento de domicilio procesal físico y ciudadanía digital.
"""

        prompt_usuario = f"""
CASO: {titulo_caso}
SUMA O PETICIÓN CONCRETA: {suma_pretendida}
NOTAS O INSTRUCCIONES COMPLEMENTARIAS: {instrucciones_especificas if instrucciones_especificas else 'Redacción forense integral y contundente.'}

HECHOS Y EVIDENCIAS PROBADAS EN EL EXPEDIENTE:
{evidencia_auditada[:4000]}
"""
        return self.llm.invocar(
            mensajes=[
                {"role": "system", "content": prompt_sistema},
                {"role": "user", "content": prompt_usuario}
            ],
            temperatura=0.08,
            max_tokens=1800,
            auto_continuar=True
        )

    @staticmethod
    def exportar_docx(texto_markdown: str, ruta_salida_docx: str):
        """Compila el texto en un archivo Word profesional listo para firma."""
        doc = Document()
        for s in doc.sections:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.3)
            s.right_margin = Inches(0.8)

        estilo_normal = doc.styles['Normal']
        estilo_normal.font.name = 'Times New Roman'
        estilo_normal.font.size = Pt(12)
        estilo_normal.font.color.rgb = RGBColor(0, 0, 0)

        for linea in texto_markdown.split("\n"):
            l = linea.strip()
            if not l:
                continue
            if l.startswith("# "):
                p = doc.add_paragraph()
                run = p.add_run(l.replace("# ", ""))
                run.bold = True
                run.font.size = Pt(13)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif l.startswith("## "):
                p = doc.add_paragraph()
                run = p.add_run(l.replace("## ", ""))
                run.bold = True
                run.font.size = Pt(12)
                p.paragraph_format.space_before = Pt(8)
            elif l.startswith("### "):
                p = doc.add_paragraph()
                run = p.add_run(l.replace("### ", ""))
                run.bold = True
            elif l.startswith("SUMA:") or "SEÑOR JUEZ" in l:
                p = doc.add_paragraph()
                run = p.add_run(l.replace("**", ""))
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(6)
            else:
                p = doc.add_paragraph()
                texto_limpio = l.replace("**", "").replace("*", "")
                p.add_run(texto_limpio)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(3)

        doc.save(ruta_salida_docx)

    @staticmethod
    def exportar_pdf(texto_markdown: str, ruta_salida_pdf: str):
        """Genera versión PDF formal del documento."""
        pdf = MarkdownPdf(toc_level=0)
        pdf.add_section(Section(texto_markdown))
        pdf.save(ruta_salida_pdf)
